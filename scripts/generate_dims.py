#!/usr/bin/env python3
"""Code generator: workbook registry -> committed source modules (plan.md §7.2, §9).

Reads the *Dayshape Reporting Service Detail* workbook and emits two checked-in
modules:

  * ``src/dayshape/dims.py``            -- typed per-entity dimension constants.
  * ``src/dayshape/reports/_catalogue.py`` -- the report registry data table.

Run with ``--check`` in CI to fail the build if the committed modules drift from
the workbook (the workbook itself is *not* shipped, so installation needs no
spreadsheet — only this generator does, at dev time).

Usage::

    python scripts/generate_dims.py            # (re)write the modules
    python scripts/generate_dims.py --check     # exit 1 if they would change
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ModuleNotFoundError:  # pragma: no cover - dev-only dependency
    print(
        "openpyxl is required for codegen. Install the dev group: "
        "pip install -e . --group dev",
        file=sys.stderr,
    )
    raise

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "specs" / "Dayshape Reporting Service Detail (v25.7.0.0).xlsx"
DIMS_OUT = ROOT / "src" / "dayshape" / "dims.py"
CATALOGUE_OUT = ROOT / "src" / "dayshape" / "reports" / "_catalogue.py"
WORKBOOK_VERSION = "25.7.0.0"

# Sheets that are not individual reports.
NON_REPORT_SHEETS = {
    "Reports Menu",
    "Version History",
    "Reporting Settings",
    "Custom Field Filters",
    "Audit Trail shared Dims-Filters",
}

# The 17 rate-limited reports, by ReportId — sourced verbatim from the API
# documentation (R2), "Rate limiting" section. A single user may run only one of
# these at a time; a second concurrent run returns HTTP 429.
RATE_LIMITED_REPORT_IDS = frozenset(
    {
        "Availability",
        "AvailabilityPlanning",
        "LogListingTask",  # Booking Audit Trail
        "TaskHoursByTime",  # Booking Hours Over Time
        "TaskListing",  # Booking Listing
        "TaskVsActualHoursByTime",  # Booking Vs Actual Hours Over Time
        "ComparativeRevenue",
        "ClashHoursByTime",  # Clash Hours Over Time
        "LogListingJob",  # Engagement Audit Trail
        "LogListingJobGroup",  # Engagement Group Audit Trail
        "LogListingWorker",  # Resource Audit Trail
        "WorkerWorkHours",  # Resource Work Hours
        "RevenueByTime",  # Revenue Over Time
        "LogListingSetting",  # Settings Audit Trail
        "UnavailabilityHoursByTime",  # Unavailability Hours Over Time
        "LogListingUnavailability",  # Unavailability Audit Trail
        "LogListingUser",  # User Audit Trail
    }
)


# --------------------------------------------------------------------------- #
# Workbook parsing
# --------------------------------------------------------------------------- #
def _rows(ws: object) -> list[tuple[object, ...]]:
    return list(ws.iter_rows(values_only=True))  # type: ignore[attr-defined]


def _label_value(rows: list[tuple[object, ...]], label: str) -> object | None:
    """Return the cell to the right of the first cell equal to ``label``."""
    for row in rows:
        for j, cell in enumerate(row):
            if cell == label:
                for k in range(j + 1, len(row)):
                    if row[k] is not None:
                        return row[k]
    return None


def _report_types() -> dict[str, str]:
    """Map report display name -> 'Listing' | 'Pivot' from the Reports Menu sheet."""
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["Reports Menu"]
    out: dict[str, str] = {}
    for row in _rows(ws):
        # Rows look like: [None, ' - Booking Listing', 'Listing', 'Link']
        name = row[1] if len(row) > 1 else None
        rtype = row[2] if len(row) > 2 else None
        if isinstance(name, str) and isinstance(rtype, str) and rtype in ("Listing", "Pivot"):
            out[name.strip(" -").strip()] = rtype
    wb.close()
    return out


def _dim_header(rows: list[tuple[object, ...]]) -> tuple[int, int] | None:
    """Find the (row index, column index) of the 'Dimension Id' header cell."""
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if cell == "Dimension Id":
                return i, j
    return None


_IDENT = re.compile(r"^[A-Za-z][A-Za-z0-9]*$")


def parse_workbook() -> tuple[list[dict[str, object]], dict[str, set[str]]]:
    """Return (reports, dims_by_namespace).

    ``reports`` is a list of ``{report_id, name, type, permissions, rate_limited}``.
    ``dims_by_namespace`` maps a namespace token (e.g. ``"Task"``) to its set of
    dimension ids.
    """
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    types = _report_types()
    reports: list[dict[str, object]] = []
    dims_by_ns: dict[str, set[str]] = {}

    for sheet in wb.sheetnames:
        if sheet in NON_REPORT_SHEETS:
            continue
        rows = _rows(wb[sheet])
        report_id = _label_value(rows, "ReportId")
        if not isinstance(report_id, str):
            continue
        report_id = report_id.strip()
        permissions = _label_value(rows, "Permissions Required")
        name = sheet.strip()
        reports.append(
            {
                "report_id": report_id,
                "name": name,
                "type": types.get(name, "Listing"),
                "permissions": (permissions or "").strip() if isinstance(permissions, str) else "",
                "rate_limited": report_id in RATE_LIMITED_REPORT_IDS,
            }
        )

        header = _dim_header(rows)
        if header is None:
            continue
        hrow, hcol = header
        for row in rows[hrow + 1 :]:
            dim = row[hcol] if len(row) > hcol else None
            if not isinstance(dim, str):
                continue
            dim = dim.strip()
            if not _IDENT.match(dim):
                # Skips blanks and the comparative-dimension JSON blobs in the
                # Comparative Revenue sheet.
                continue
            ns, _const = split_namespace(dim)
            dims_by_ns.setdefault(ns, set()).add(dim)

    wb.close()
    # Deduplicate reports by id (a report id only appears once), keep stable order.
    seen: set[str] = set()
    unique_reports: list[dict[str, object]] = []
    for r in sorted(reports, key=lambda r: str(r["report_id"])):
        rid = str(r["report_id"])
        if rid in seen:
            continue
        seen.add(rid)
        unique_reports.append(r)
    return unique_reports, dims_by_ns


# --------------------------------------------------------------------------- #
# Naming
# --------------------------------------------------------------------------- #
_LEAD_WORD = re.compile(r"[A-Za-z][a-z0-9]*")


def split_namespace(dim: str) -> tuple[str, str]:
    """Split a dimension id into (namespace token, remainder).

    ``JobGroup*`` is treated as a single token so it does not collapse into the
    ``Job`` namespace. Otherwise the leading lower-case-tailed word is the token.
    """
    if dim.lower().startswith("jobgroup"):
        return "JobGroup", dim[len("JobGroup") :]
    m = _LEAD_WORD.match(dim)
    word = m.group(0) if m else dim
    token = word[0].upper() + word[1:].lower()
    return token, dim[len(word) :]


def const_name(remainder: str, full: str) -> str:
    """Convert the remainder after the namespace token into UPPER_SNAKE."""
    source = remainder or full
    # camelCase / PascalCase -> snake
    snake = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", source)
    snake = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", snake)
    snake = snake.replace("-", "_").replace(" ", "_")
    out = snake.upper().strip("_")
    if not out or out[0].isdigit():
        out = "D_" + out
    return out


# --------------------------------------------------------------------------- #
# Emission
# --------------------------------------------------------------------------- #
BANNER = (
    "# AUTO-GENERATED by scripts/generate_dims.py — DO NOT EDIT BY HAND.\n"
    f"# Source: Dayshape Reporting Service Detail workbook v{WORKBOOK_VERSION}.\n"
    "# Regenerate with: python scripts/generate_dims.py\n"
)


def render_dims(dims_by_ns: dict[str, set[str]]) -> str:
    lines: list[str] = []
    lines.append('"""Typed dimension constants generated from the workbook (plan.md §7.2).')
    lines.append("")
    lines.append("Each ``*Dim`` namespace groups the dimension ids of one entity. ``Dim`` is a")
    lines.append("``str`` subclass, so a constant is usable anywhere a raw id string is accepted")
    lines.append('(``"TaskId" == TaskDim.ID``) and serialises for free. ``.asc()``/``.desc()``')
    lines.append("return a sorted :class:`~dayshape.reports.query.Dimension`.")
    lines.append('"""')
    lines.append("")
    lines.append("from __future__ import annotations")
    lines.append("")
    lines.append("from typing import TYPE_CHECKING")
    lines.append("")
    lines.append("if TYPE_CHECKING:")
    lines.append("    from .reports.query import Dimension")
    lines.append("")
    lines.append("")
    lines.append("class Dim(str):")
    lines.append('    """A dimension id. ``str`` subclass with ordering sugar."""')
    lines.append("")
    lines.append("    __slots__ = ()")
    lines.append("")
    lines.append('    def asc(self) -> "Dimension":')
    lines.append('        """This dimension, sorted ascending."""')
    lines.append("        from .reports.query import Dimension, Sort")
    lines.append("")
    lines.append("        return Dimension(self, order=Sort.ASC)")
    lines.append("")
    lines.append('    def desc(self) -> "Dimension":')
    lines.append('        """This dimension, sorted descending."""')
    lines.append("        from .reports.query import Dimension, Sort")
    lines.append("")
    lines.append("        return Dimension(self, order=Sort.DESC)")
    lines.append("")

    all_names: list[str] = []
    for ns in sorted(dims_by_ns):
        class_name = f"{ns}Dim"
        all_names.append(class_name)
        lines.append("")
        lines.append(f"class {class_name}:")
        lines.append(f'    """Dimensions of the {ns} entity."""')
        lines.append("")
        # Build const -> dim, resolving collisions deterministically.
        consts: dict[str, str] = {}
        for dim in sorted(dims_by_ns[ns]):
            _, remainder = split_namespace(dim)
            name = const_name(remainder, dim)
            if name in consts and consts[name] != dim:
                name = const_name(dim, dim)  # disambiguate with the full id
                suffix = 2
                base = name
                while name in consts and consts[name] != dim:
                    name = f"{base}_{suffix}"
                    suffix += 1
            consts[name] = dim
        for name in sorted(consts):
            lines.append(f'    {name} = Dim("{consts[name]}")')
        lines.append("")

    lines.append("")
    lines.append("__all__ = [")
    lines.append('    "Dim",')
    for class_name in all_names:
        lines.append(f'    "{class_name}",')
    lines.append("]")
    lines.append("")
    return BANNER + "\n" + "\n".join(lines)


def render_catalogue(reports: list[dict[str, object]]) -> str:
    lines: list[str] = []
    lines.append('"""Generated report registry data (plan.md §0.4, §3.5).')
    lines.append("")
    lines.append("Consumed by :mod:`dayshape.reports.registry`. Permission strings are kept")
    lines.append("verbatim from the workbook, including the known irregularities flagged in")
    lines.append("plan.md §11.2 R-1 (e.g. ``Report_UnitListing_Read``).")
    lines.append('"""')
    lines.append("")
    lines.append("from __future__ import annotations")
    lines.append("")
    lines.append("# report_id -> (display_name, report_type, permissions, rate_limited)")
    lines.append("CATALOGUE: dict[str, tuple[str, str, str, bool]] = {")
    for r in reports:
        rid = r["report_id"]
        name = r["name"]
        rtype = r["type"]
        perms = str(r["permissions"]).replace("\n", " ").replace("\r", " ").strip()
        perms = re.sub(r"\s+", " ", perms)
        rl = r["rate_limited"]
        lines.append(f"    {rid!r}: ({name!r}, {rtype!r}, {perms!r}, {rl!r}),")
    lines.append("}")
    lines.append("")
    return BANNER + "\n" + "\n".join(lines)


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #
def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Exit non-zero if the committed modules would change.",
    )
    args = parser.parse_args()

    reports, dims_by_ns = parse_workbook()
    dims_src = render_dims(dims_by_ns)
    catalogue_src = render_catalogue(reports)

    targets = [(DIMS_OUT, dims_src), (CATALOGUE_OUT, catalogue_src)]

    if args.check:
        drift = False
        for path, content in targets:
            current = path.read_text(encoding="utf-8") if path.exists() else ""
            if current != content:
                print(f"DRIFT: {path.relative_to(ROOT)} is out of date.", file=sys.stderr)
                drift = True
        if drift:
            print("Run: python scripts/generate_dims.py", file=sys.stderr)
            return 1
        print("dims.py and _catalogue.py are up to date.")
        return 0

    for path, content in targets:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        print(f"wrote {path.relative_to(ROOT)}")
    n_dims = sum(len(v) for v in dims_by_ns.values())
    print(f"  {len(reports)} reports, {n_dims} dimensions, {len(dims_by_ns)} namespaces")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
